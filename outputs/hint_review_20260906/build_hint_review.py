import json
import sys
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo

root = Path(r'C:\Users\User\AndroidStudioProjects\MysteryLane')
rows = []
for line in sys.stdin:
    line = line.strip()
    if line == '__END__':
        break
    if line:
        rows.extend(json.loads(line))

h1_counts = Counter(str(row.get('hint_1') or '').strip().lower() for row in rows)
h2_counts = Counter(str(row.get('hint_2') or '').strip().lower() for row in rows)
data = []
counts = defaultdict(lambda: [0, 0, 0])
for row in rows:
    destination = row.get('destination_name') or (row.get('blind_box_destinations') or {}).get('name') or 'Malaysia General Knowledge'
    puzzle_type = row['puzzle_type']
    hint1 = str(row.get('hint_1') or '').strip()
    hint2 = str(row.get('hint_2') or '').strip()
    index = {'Guess the Word': 0, 'Multiple Choice Question': 1, 'True or False': 2}[puzzle_type]
    counts[destination][index] += 1
    data.append([
        destination, puzzle_type, row['question_text'], row.get('display_box_content') or '',
        row['correct_answer'], hint1, hint2,
        'Unique' if h1_counts[hint1.lower()] == 1 else 'Repeated',
        'Unique' if h2_counts[hint2.lower()] == 1 else 'Repeated', row['puzzle_id'],
    ])
data.sort(key=lambda row: (row[0].lower(), row[1], row[2].lower()))

wb = Workbook()
summary = wb.active
summary.title = 'Destination summary'
questions = wb.create_sheet('Question hints')
summary.sheet_view.showGridLines = False
questions.sheet_view.showGridLines = False

summary.merge_cells('A1:E1')
summary['A1'] = 'Puzzle hint inventory'
summary.merge_cells('A2:E2')
summary['A2'] = 'Active questions in the three playable categories. No database values were changed.'
summary.append([])
summary.append(['Destination', 'Scrambled Anagrams', 'Multiple Choice Trivia', 'True or False', 'Total'])
for name in sorted(counts, key=str.lower):
    values = counts[name]
    summary.append([name, *values, sum(values)])

headers = ['Destination', 'Category', 'Question', 'Answer to check', 'Correct answer', 'Hint 1', 'Hint 2', 'Hint 1 status', 'Hint 2 status', 'Puzzle ID']
questions.append(headers)
for row in data:
    questions.append(row)

header_fill = PatternFill('solid', fgColor='087FAC')
for sheet, header_row in ((summary, 4), (questions, 1)):
    for cell in sheet[header_row]:
        cell.fill = header_fill
        cell.font = Font(name='Arial', size=10, bold=True, color='FFFFFF')
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    for row in sheet.iter_rows():
        for cell in row:
            if cell.row != header_row:
                cell.font = Font(name='Arial', size=10, color='172033')
                cell.alignment = Alignment(vertical='top', wrap_text=True)

summary['A1'].font = Font(name='Arial', size=16, bold=True, color='172033')
summary['A2'].font = Font(name='Arial', size=10, italic=True, color='64748B')
summary.column_dimensions['A'].width = 44
for col in 'BCDE': summary.column_dimensions[col].width = 20
summary.freeze_panes = 'A5'

widths = [34, 24, 55, 38, 28, 42, 42, 15, 15, 36]
for index, width in enumerate(widths, 1):
    questions.column_dimensions[chr(64 + index)].width = width
questions.freeze_panes = 'C2'
table = Table(displayName='QuestionHintsTable', ref=f'A1:J{len(data) + 1}')
table.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True, showColumnStripes=False)
questions.add_table(table)

output = root / 'outputs/hint_review_20260906/puzzle_question_hints.xlsx'
wb.save(output)

check = load_workbook(output, read_only=False, data_only=False)
assert check.sheetnames == ['Destination summary', 'Question hints']
assert check['Question hints'].max_row == len(data) + 1
assert check['Destination summary'].max_row == len(counts) + 4
print(json.dumps({'output': str(output), 'questions': len(data), 'destinations': len(counts)}))
